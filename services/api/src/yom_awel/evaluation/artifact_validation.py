"""Signature-first inspection of untrusted uploads, then a bounded read into a text table.

CSV and XLSX both become the same ``Table`` of strings, so grading never depends on the
upload format.
"""

import csv
import io
import re
import warnings
import zipfile
from collections.abc import Iterable, Mapping
from dataclasses import dataclass
from datetime import datetime, time
from pathlib import PurePosixPath
from xml.etree.ElementTree import Element

from defusedxml.ElementTree import iterparse  # type: ignore[import-untyped]
from openpyxl import load_workbook
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

from yom_awel.domain.contracts import ArtifactRef
from yom_awel.domain.errors import DomainError
from yom_awel.evaluation.task_package import ArtifactLimits

ZIP_SIGNATURE = b"PK\x03\x04"
OLE_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
UTF8_BOM = "﻿"
REQUIRED_XLSX_ENTRIES = frozenset({"[Content_Types].xml", "xl/workbook.xml"})
SHEET_DIRECTORIES = frozenset({PurePosixPath("xl/worksheets"), PurePosixPath("xl/chartsheets")})
# Macros, Excel 4 macro sheets and links to other workbooks are never read.
PROHIBITED_XLSX_PARTS = ("xl/vbaProject.bin", "xl/macrosheets/", "xl/externalLinks/")
# The v1 vocabulary has no row or column code; a table past either bound exceeds the
# safe processing limit.
TABLE_LIMIT_CODE = "expanded_size_exceeded"
CELL_REFERENCE = re.compile(r"^([A-Z]+)([1-9][0-9]*)$")

# Rejection vocabulary v1, verbatim from task_packages/clean-sales/1/learning-objectives.yaml.
REJECTION_MESSAGES: Mapping[str, tuple[str, str]] = {
    "unsupported_type": (
        "صيغة الملف غير مدعومة. يرجى رفع ملف CSV أو XLSX صالح.",
        "Unsupported file format. Please upload a valid CSV or XLSX spreadsheet.",
    ),
    "artifact_too_large": (
        "حجم الملف يتجاوز الحد الأقصى المسموح (5 ميجابايت).",
        "File size exceeds the 5 MiB limit.",
    ),
    "mime_mismatch": (
        "محتوى الملف لا يطابق صيغة CSV أو XLSX المعلنة.",
        "File content does not match its declared CSV or XLSX format.",
    ),
    "expanded_size_exceeded": (
        "محتوى ملف العمل بعد فك الضغط يتجاوز حد المعالجة الآمن.",
        "Expanded workbook content exceeds the safe processing limit.",
    ),
    "sheet_limit_exceeded": (
        "يجب أن يحتوي ملف العمل على ورقة العمل النشطة المدعومة فقط.",
        "Workbook must contain only the supported active worksheet.",
    ),
    "artifact_unreadable": (
        "تعذر قراءة الملف كجدول بيانات صالح (مثل بنية تالفة، أو ماكرو مفعّل، أو ملف محمي بكلمة سر).",
        (
            "File cannot be parsed as a tabular spreadsheet "
            "(e.g. corrupt structure, macro workbook, or password protected)."
        ),
    ),
    "missing_columns": (
        "أعمدة أساسية مطلوبة مفقودة من الملف.",
        "Required column headers are missing.",
    ),
    "duplicate_columns": (
        "تم العثور على أسماء أعمدة مكررة في صف العناوين.",
        "Duplicate column names detected in the header row.",
    ),
    "too_few_rows": (
        "الملف المنظف يحتوي على أقل من الحد الأدنى المطلوب (40 صفاً).",
        "Cleaned dataset contains fewer than the required 40 rows.",
    ),
}

Row = tuple[str, ...]


class ArtifactRejected(DomainError):
    """The learner's file cannot be graded; the attempt ends with a retryable result."""

    def __init__(self, code: str) -> None:
        message_ar, message_en = REJECTION_MESSAGES[code]
        super().__init__(code, message_en, message_ar=message_ar)


class ArtifactSizeMismatch(DomainError):
    """Storage handed over bytes that disagree with the upload record; nothing is graded."""

    def __init__(self) -> None:
        super().__init__(
            "artifact_size_mismatch", "Artifact content does not match its recorded size"
        )


@dataclass(frozen=True)
class Table:
    rows: tuple[Row, ...]


def inspect_artifact(artifact: ArtifactRef, limits: ArtifactLimits) -> Table:
    """Inspect ``artifact.content``, the bytes the application loaded for evaluation."""

    content = artifact.content
    if max(len(content), artifact.size_bytes) > limits.max_bytes:
        raise ArtifactRejected("artifact_too_large")
    if len(content) != artifact.size_bytes:
        raise ArtifactSizeMismatch()
    extension = PurePosixPath(artifact.filename).suffix.lower().lstrip(".")
    if extension == "xlsx" and "xlsx" in limits.extensions:
        return _bounded(_read_xlsx(content, limits), limits)
    if extension == "csv" and "csv" in limits.extensions:
        return _bounded(_read_csv(content), limits)
    raise ArtifactRejected("unsupported_type")


def _bounded(rows: Iterable[Row], limits: ArtifactLimits) -> Table:
    table = tuple(rows)
    if len(table) > limits.max_rows + 1 or any(len(row) > limits.max_columns for row in table):
        raise ArtifactRejected(TABLE_LIMIT_CODE)
    return Table(rows=table)


def _read_csv(content: bytes) -> list[Row]:
    if content.startswith((ZIP_SIGNATURE, OLE_SIGNATURE)) or b"\x00" in content:
        raise ArtifactRejected("mime_mismatch")
    try:
        text = content.decode("utf-8").removeprefix(UTF8_BOM)
        return [tuple(row) for row in csv.reader(io.StringIO(text, newline=""))]
    except (UnicodeDecodeError, csv.Error) as error:
        raise ArtifactRejected("artifact_unreadable") from error


def _read_xlsx(content: bytes, limits: ArtifactLimits) -> list[Row]:
    # Password-protected workbooks are OLE containers, not ZIP packages.
    if content.startswith(OLE_SIGNATURE):
        raise ArtifactRejected("artifact_unreadable")
    if not content.startswith(ZIP_SIGNATURE):
        raise ArtifactRejected("mime_mismatch")
    _check_package(content, limits)
    try:
        _reject_nonempty_xlsx_overflow(content, limits)
        values = _active_sheet_values(content, limits)
    except ArtifactRejected:
        raise
    except Exception as error:  # Any parser failure on untrusted input is unreadable.
        raise ArtifactRejected("artifact_unreadable") from error
    rows = [_text_row(row) for row in values]
    while rows and not rows[-1]:
        rows.pop()
    return rows


def _reject_nonempty_xlsx_overflow(content: bytes, limits: ArtifactLimits) -> None:
    """Reject real cells past the bounded reader's window before reading their values.

    ``openpyxl`` stops at the requested row/column window. A sparse cell beyond
    an empty boundary would otherwise disappear from the table and be graded as
    though it were never uploaded. The package has one worksheet, so scanning
    its cell references is bounded by the already-checked expanded ZIP size.
    """

    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        sheets = [name for name in archive.namelist() if _is_sheet(name)]
        if len(sheets) != 1:
            raise ArtifactRejected("artifact_unreadable")
        with archive.open(sheets[0]) as source:
            for _, element in iterparse(source, events=("end",)):
                if element.tag.endswith("}row"):
                    element.clear()
                    continue
                if not element.tag.endswith("}c"):
                    continue
                if not _cell_has_value(element):
                    element.clear()
                    continue
                reference = element.attrib.get("r")
                match = CELL_REFERENCE.fullmatch(reference or "")
                if match is None:
                    raise ArtifactRejected("artifact_unreadable")
                column = _column_number(match.group(1))
                row = int(match.group(2))
                if row > limits.max_rows + 1 or column > limits.max_columns:
                    raise ArtifactRejected(TABLE_LIMIT_CODE)
                element.clear()


def _cell_has_value(element: Element) -> bool:
    return any(
        child.tag.endswith(("}v", "}f", "}is")) and (child.text or list(child))
        for child in element.iter()
    )


def _column_number(letters: str) -> int:
    value = 0
    for letter in letters:
        value = value * 26 + ord(letter) - ord("A") + 1
    return value


def _check_package(content: bytes, limits: ArtifactLimits) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
    except (zipfile.BadZipFile, ValueError) as error:
        raise ArtifactRejected("artifact_unreadable") from error
    names = {entry.filename for entry in entries}
    if not REQUIRED_XLSX_ENTRIES <= names or not all(map(_is_safe_entry, names)):
        raise ArtifactRejected("artifact_unreadable")
    if any(name.startswith(PROHIBITED_XLSX_PARTS) for name in names):
        raise ArtifactRejected("artifact_unreadable")
    # zipfile stops at each entry's declared size, so this bounds decompression.
    if sum(entry.file_size for entry in entries) > limits.max_expanded_bytes:
        raise ArtifactRejected("expanded_size_exceeded")
    sheets = sum(1 for name in names if _is_sheet(name))
    if sheets == 0:
        raise ArtifactRejected("artifact_unreadable")
    if sheets > limits.max_sheets:
        raise ArtifactRejected("sheet_limit_exceeded")


def _active_sheet_values(content: bytes, limits: ArtifactLimits) -> list[tuple[object, ...]]:
    """Read one row and one column past each limit, so ``_bounded`` can reject overflow.

    Cells further out are never read, which bounds work on a sheet with forged row numbers.
    Formulas are read as their text, never as cached results.
    """

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
        try:
            sheet = workbook.active
            if not isinstance(sheet, ReadOnlyWorksheet):
                raise TypeError("the active sheet is not a worksheet")
            return list(
                sheet.iter_rows(
                    max_row=limits.max_rows + 2,
                    max_col=limits.max_columns + 1,
                    values_only=True,
                )
            )
        finally:
            workbook.close()


def _cell_text(value: object) -> str:
    """The text a CSV export of the cell would hold; a date cell at midnight is its ISO date."""

    match value:
        case None:
            return ""
        case bool():
            return "TRUE" if value else "FALSE"
        case datetime() if value.time() == time():
            return value.date().isoformat()
        case float():
            return repr(value)
        case _:
            return str(value)


def _text_row(values: Iterable[object]) -> Row:
    cells = [_cell_text(value) for value in values]
    while cells and not cells[-1]:
        cells.pop()
    return tuple(cells)


def _is_safe_entry(name: str) -> bool:
    path = PurePosixPath(name)
    return not path.is_absolute() and ".." not in path.parts and "\\" not in name


def _is_sheet(name: str) -> bool:
    path = PurePosixPath(name)
    return path.parent in SHEET_DIRECTORIES and path.suffix == ".xml"
