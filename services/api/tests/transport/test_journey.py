import hashlib
import io
import zipfile

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from tests.transport.support import (
    CLEAN_CSV,
    CLEAN_XLSX,
    CSRF,
    DIRTY_CSV,
    DUPLICATES_CSV,
    TASK_ID,
    XLSX_MIME,
    browser,
    onboard,
    settings_for,
    start,
    submit,
)
from yom_awel.evaluation.catalog import csv_to_xlsx
from yom_awel.transport.app import create_app


@pytest.fixture
def settings(tmp_path):
    return settings_for(tmp_path)


@pytest.fixture
def client(settings):
    with TestClient(create_app(settings), headers=CSRF) as client:
        yield client


def check_codes(outcome):
    return {
        check["check_id"]: check["diagnostic_code"] for check in outcome["evaluation"]["checks"]
    }


def test_health_runtime_and_protected_routes(client):
    assert client.get("/api/v1/health").json() == {"status": "ok"}
    assert client.get("/api/v1/runtime").json() == {
        "mode": "local",
        "feedback_provider": "deterministic",
    }
    for path in ("/api/v1/tasks", "/api/v1/tasks/current", f"/api/v1/tasks/{TASK_ID}"):
        assert client.get(path).status_code == 401


def test_state_changing_requests_need_the_csrf_header(settings):
    with TestClient(create_app(settings)) as bare:
        response = bare.post("/api/v1/auth/session")
        assert response.status_code == 403
        assert "set-cookie" not in response.headers
        assert bare.post("/api/v1/auth/session", headers=CSRF).status_code == 200


def test_session_cookie_is_http_only_lax_and_reused(client):
    response = client.post("/api/v1/auth/session")
    cookie = response.headers["set-cookie"].lower()
    assert "yom_session=" in cookie
    assert "httponly" in cookie
    assert "samesite=lax" in cookie
    assert "path=/" in cookie
    assert "secure" not in cookie  # local HTTP; cloud mode sets Secure
    assert "access_token" not in response.text
    again = client.post("/api/v1/auth/session")
    assert "set-cookie" not in again.headers
    logout = client.post("/api/v1/auth/logout")
    assert logout.status_code == 204
    assert 'yom_session=""' in logout.headers["set-cookie"]


def test_catalog_start_detail_and_downloads(client):
    onboard(client)
    [summary] = client.get("/api/v1/tasks").json()["tasks"]
    assert summary["task_id"] == TASK_ID
    assert summary["status"] == "available"
    assert (summary["pass_threshold"], summary["points_total"]) == (75, 100)
    assert summary["title_en"] == "Practical Assignment: Cleaning Daily Sales Transactions"
    assert summary["title_ar"].startswith("تكليف عملي")
    assert client.get("/api/v1/tasks/current").json() == {"status": "READY", "task": None}

    task = start(client)
    assert task["task_version_id"] == summary["task_version_id"]
    assert task["instructions_en"].startswith("# Practical Assignment")
    assert start(client) == task
    assert client.get("/api/v1/tasks").json()["tasks"][0]["status"] == "in_progress"

    detail = client.get(f"/api/v1/tasks/{TASK_ID}").json()
    assert [(c["check_id"], c["points"], c["critical"]) for c in detail["checks"]] == [
        ("unique_orders", 25, True),
        ("standard_dates", 25, False),
        ("valid_numeric_values", 25, False),
        ("complete_customer_records", 25, False),
    ]
    assert detail["formats"] == ["csv", "xlsx"]
    assert detail["hints_ar"] and detail["hints_en"] and detail["brief_ar"]
    assert client.get("/api/v1/tasks/unknown").status_code == 404

    csv_file = client.get(f"/api/v1/tasks/{TASK_ID}/dataset")
    assert csv_file.content == DIRTY_CSV
    assert 'filename="sales_dirty.csv"' in csv_file.headers["content-disposition"]
    xlsx_file = client.get(f"/api/v1/tasks/{TASK_ID}/dataset", params={"format": "xlsx"})
    assert xlsx_file.headers["content-type"] == XLSX_MIME
    rows = list(load_workbook(io.BytesIO(xlsx_file.content)).active.iter_rows(values_only=True))
    assert rows[0][0] == "order_id"
    assert len(rows) == len(DIRTY_CSV.decode().strip().splitlines())
    assert client.get(f"/api/v1/tasks/{TASK_ID}/dataset", params={"format": "pdf"}).status_code == (
        422
    )


def test_real_evaluator_journey_fail_critical_retry_pass_and_restart(client, settings):
    onboard(client)
    start(client)

    dirty, _, _ = submit(client, DIRTY_CSV, "dirty")
    assert (dirty["evaluation"]["score"], dirty["evaluation"]["passed"]) == (0, False)
    assert dirty["learner_status"] == "NEEDS_RETRY"
    assert set(check_codes(dirty).values()) == {
        "unique_orders_failed",
        "standard_dates_failed",
        "valid_numeric_values_failed",
        "complete_customer_records_failed",
    }

    critical, _, _ = submit(client, DUPLICATES_CSV, "duplicates")
    assert (critical["evaluation"]["score"], critical["evaluation"]["passed"]) == (75, False)
    assert check_codes(critical)["unique_orders"] == "unique_orders_failed"
    assert critical["learner_status"] == "NEEDS_RETRY"

    passed, body, keyed = submit(client, CLEAN_XLSX, "clean", "sales.xlsx", XLSX_MIME)
    assert (passed["evaluation"]["score"], passed["evaluation"]["passed"]) == (100, True)
    assert passed["learner_status"] == "TASK_COMPLETED"
    assert passed["attempt_number"] == 3
    assert passed["feedback"]["language"] == "ar-EG"
    assert client.post("/api/v1/submissions", json=body, headers=keyed).json() == passed
    assert client.get("/api/v1/submissions/" + passed["submission_id"], headers=keyed).json() == (
        passed
    )

    assert [a["attempt_number"] for a in client.get("/api/v1/attempts").json()] == [1, 2, 3]
    assert client.get("/api/v1/skills").json()["skills"][0] == {
        "skill_id": "data_cleaning",
        "score": 100,
    }
    assert client.get("/api/v1/tasks").json()["tasks"][0]["status"] == "completed"

    with TestClient(create_app(settings), headers=CSRF, cookies=client.cookies) as restarted:
        assert len(restarted.get("/api/v1/attempts").json()) == 3
        assert restarted.get("/api/v1/tasks/current").json()["status"] == "TASK_COMPLETED"


def test_csv_and_xlsx_of_the_same_table_score_the_same(client):
    onboard(client)
    start(client)
    from_csv, _, _ = submit(client, DUPLICATES_CSV, "csv")
    from_xlsx, _, _ = submit(client, csv_to_xlsx(DUPLICATES_CSV), "xlsx", "sales.xlsx", XLSX_MIME)
    assert from_csv["evaluation"]["checks"] == from_xlsx["evaluation"]["checks"]
    assert from_xlsx["evaluation"]["score"] == 75


def test_language_preference_and_feedback_in_both_languages(client):
    onboard(client, language="en")
    start(client)
    first, _, _ = submit(client, DUPLICATES_CSV, "first")
    assert first["feedback"]["language"] == "en"

    path = f"/api/v1/submissions/{first['submission_id']}/feedback"
    assert client.get(path, params={"language": "en"}).json() == first["feedback"]
    arabic = client.get(path, params={"language": "ar-EG"}).json()
    assert arabic["language"] == "ar-EG"
    assert arabic["persona_id"] == "tarek"
    assert "القرار:" in arabic["feedback_text"]
    assert client.get(path, params={"language": "fr"}).status_code == 422

    changed = client.put("/api/v1/learners/me/language", json={"preferred_language": "ar-EG"})
    assert changed.json()["preferred_language"] == "ar-EG"
    second, _, _ = submit(client, CLEAN_CSV, "second")
    assert second["feedback"]["language"] == "ar-EG"
    assert client.put(
        "/api/v1/learners/me/language", json={"preferred_language": "de"}
    ).status_code == (422)


def test_learners_cannot_see_each_other(client):
    onboard(client)
    start(client)
    outcome, body, _ = submit(client, DIRTY_CSV, "private")
    other = browser(client.app)
    onboard(other, "منى")
    start(other)
    assert (
        other.get(
            "/api/v1/submissions/" + outcome["submission_id"],
            headers={"Idempotency-Key": "private"},
        ).status_code
        == 404
    )
    assert (
        other.get(
            f"/api/v1/submissions/{outcome['submission_id']}/feedback?language=en"
        ).status_code
        == 404
    )
    stolen = other.post("/api/v1/submissions", json=body, headers={"Idempotency-Key": "stolen"})
    assert stolen.status_code in (403, 404, 422)
    injected = client.post(
        "/api/v1/learners/onboard", json={"display_name": "a", "learner_id": "injected"}
    )
    assert injected.status_code == 422


def test_starting_a_task_needs_an_onboarded_ready_learner(client):
    assert client.post("/api/v1/auth/session").status_code == 200
    assert client.post(f"/api/v1/tasks/{TASK_ID}/start").status_code == 404
    onboard(client)
    assert client.post("/api/v1/tasks/unknown/start").status_code == 404


@pytest.mark.parametrize(
    "filename,mime",
    [("payload.exe", "application/octet-stream"), ("../x.csv", "text/csv"), ("x.csv", "text/html")],
)
def test_invalid_file_metadata(client, filename, mime):
    onboard(client)
    start(client)
    response = client.post(
        "/api/v1/artifacts/upload-authorization",
        json={
            "filename": filename,
            "content_type": mime,
            "size_bytes": 2,
            "artifact_sha256": "a" * 64,
        },
    )
    assert response.status_code == 415


def test_upload_hash_size_token_and_owner(client):
    onboard(client)
    start(client)
    content = CLEAN_CSV
    metadata = {
        "filename": "x.csv",
        "content_type": "text/csv",
        "size_bytes": len(content),
        "artifact_sha256": hashlib.sha256(content).hexdigest(),
    }
    auth = client.post("/api/v1/artifacts/upload-authorization", json=metadata).json()
    assert client.put(auth["upload_url"], content=content).status_code == 403
    assert (
        client.put(
            auth["upload_url"], headers=auth["headers"], content=b"x" * len(content)
        ).status_code
        == 422
    )
    other = browser(client.app)
    onboard(other, "other")
    assert (
        other.put(auth["upload_url"], headers=auth["headers"], content=content).status_code == 403
    )
    for _ in range(2):
        assert (
            client.put(auth["upload_url"], headers=auth["headers"], content=content).status_code
            == 200
        )


def test_cors_and_size_limit(client):
    for origin, status in [("http://localhost:3000", 200), ("https://evil.test", 400)]:
        response = client.options(
            "/api/v1/tasks/current",
            headers={"Origin": origin, "Access-Control-Request-Method": "GET"},
        )
        assert response.status_code == status
    allowed = client.options(
        "/api/v1/tasks/current",
        headers={"Origin": "http://localhost:3000", "Access-Control-Request-Method": "GET"},
    )
    assert allowed.headers["access-control-allow-origin"] == "http://localhost:3000"
    assert client.post("/api/v1/learners/onboard", content=b"x" * 65537).status_code == 413
    assert (
        client.put("/api/v1/artifacts/no/content", content=b"x" * (5 * 1024 * 1024 + 1)).status_code
        == 413
    )


def test_xlsx_download_is_a_single_sheet_workbook(client):
    onboard(client)
    start(client)
    content = client.get(f"/api/v1/tasks/{TASK_ID}/dataset", params={"format": "xlsx"}).content
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        sheets = [name for name in archive.namelist() if name.startswith("xl/worksheets/")]
    assert sheets == ["xl/worksheets/sheet1.xml"]
