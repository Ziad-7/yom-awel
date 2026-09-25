"""The presenter kit: every demo upload graded by the real sales-cleaning evaluator."""

import asyncio
import hashlib
import io
import zipfile
from dataclasses import dataclass
from pathlib import Path
from uuid import UUID

import pytest
from openpyxl import load_workbook
from yom_awel.domain.contracts import ArtifactRef, EvaluationResult, TaskVersion
from yom_awel.evaluation.clean_sales_dataset import COLUMNS, LEARNER_SEED, generate
from yom_awel.evaluation.sales_cleaning import SalesCleaningEvaluator

from demo.generate_demo_files import DEMO_FILES, OUTPUT_DIR, build_all, write_all

ROOT = Path(__file__).resolve().parents[3]
PACKAGE_ROOT = ROOT / "task_packages" / "clean-sales" / "1"
TASK_VERSION = TaskVersion.model_validate_json(
    (ROOT / "contracts" / "fixtures" / "task-version-clean-sales.json").read_text(
        "utf-8"
    )
)
EVALUATOR = SalesCleaningEvaluator.from_directory(PACKAGE_ROOT)
ALL_CHECKS = [
    "unique_orders",
    "standard_dates",
    "valid_numeric_values",
    "complete_customer_records",
]


@dataclass(frozen=True)
class Expected:
    score: int
    passed: bool
    failed_checks: list[str]
    error_code: str | None = None


EXPECTED = {
    "sales_cleaned.csv": Expected(100, True, []),
    "sales_cleaned.xlsx": Expected(100, True, []),
    "sales_retry_duplicates.csv": Expected(75, False, ["unique_orders"]),
    "sales_retry_half.xlsx": Expected(
        50, False, ["standard_dates", "complete_customer_records"]
    ),
    "sales_rejected_missing_columns.csv": Expected(
        0, False, ALL_CHECKS, "missing_columns"
    ),
}


def grade(name: str, content: bytes) -> EvaluationResult:
    artifact = ArtifactRef(
        artifact_id=UUID("00000000-0000-0000-0000-0000000000de"),
        filename=name,
        size_bytes=len(content),
        sha256=hashlib.sha256(content).hexdigest(),
        content=content,
    )
    return asyncio.run(EVALUATOR.evaluate(TASK_VERSION, artifact))


def test_every_demo_file_has_an_expectation() -> None:
    assert set(EXPECTED) == set(DEMO_FILES)


@pytest.mark.parametrize("name", sorted(EXPECTED))
def test_committed_demo_file_grades_as_scripted(name: str) -> None:
    expected = EXPECTED[name]

    result = grade(name, (OUTPUT_DIR / name).read_bytes())

    assert result.score == expected.score
    assert result.passed is expected.passed
    assert [check.check_id for check in result.checks if not check.passed] == (
        expected.failed_checks
    )
    assert [error.code for error in result.errors] == (
        [expected.error_code] if expected.error_code else []
    )


def test_learner_dirty_file_fails_every_check() -> None:
    result = grade(
        "sales_dirty.csv", (PACKAGE_ROOT / "data" / "sales_dirty.csv").read_bytes()
    )

    assert (result.score, result.passed, result.errors) == (0, False, [])


@pytest.mark.parametrize("name", sorted(EXPECTED))
def test_committed_demo_file_equals_a_fresh_regeneration(name: str) -> None:
    assert (OUTPUT_DIR / name).read_bytes() == build_all(generate(LEARNER_SEED))[name]


@pytest.mark.parametrize("name", ["sales_cleaned.xlsx", "sales_retry_half.xlsx"])
def test_workbook_has_one_sheet_and_fixed_properties(name: str) -> None:
    content = (OUTPUT_DIR / name).read_bytes()
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        assert all(info.create_system == 3 for info in archive.infolist())
    book = load_workbook(io.BytesIO(content), read_only=True)
    try:
        header = next(book.active.iter_rows(max_row=1, values_only=True))  # type: ignore[union-attr]
        assert book.sheetnames == ["sales"]
        assert header == COLUMNS
        assert book.properties.creator == "Yom Awel demo kit"
        assert str(book.properties.modified) == "2026-09-20 00:00:00"
    finally:
        book.close()


def test_write_all_writes_every_file(tmp_path: Path) -> None:
    files = build_all(generate(LEARNER_SEED))

    paths = write_all(tmp_path / "kit", files)

    assert {path.name: path.read_bytes() for path in paths} == files
